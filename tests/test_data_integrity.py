"""发布数据的连接键、缺失语义、聚合量与工作簿一致性检查。"""
from pathlib import Path
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def read(name):
    return pd.read_csv(ROOT / name, keep_default_na=False)


def identifiers(series):
    return {part.strip() for cell in series for part in str(cell).split(';') if part.strip()}


class DataIntegrityTests(unittest.TestCase):
    def test_trait_matrix_and_traceability(self):
        taxa = read('data/traits/taxa.csv')
        schema = read('data/traits/trait_schema.csv')
        wide = read('data/traits/trait_matrix.csv').set_index('taxon_id')
        values = read('data/traits/trait_values.csv')
        self.assertEqual(wide.shape, (37, 37))
        self.assertEqual(set(wide.index), set(taxa.taxon_id))
        self.assertEqual(set(wide.columns), set(schema.output_column))
        self.assertFalse(values.duplicated(['taxon_id', 'output_column']).any())
        reconstructed = values.pivot(index='taxon_id', columns='output_column', values='value')
        reconstructed = reconstructed.loc[wide.index, wide.columns]
        a = wide.replace('', np.nan).astype(float).to_numpy()
        b = reconstructed.replace('', np.nan).astype(float).to_numpy()
        np.testing.assert_allclose(a, b, equal_nan=True)
        self.assertEqual(int(np.isnan(a).sum()), 398)
        self.assertEqual(int(np.isfinite(a).sum()), 971)

    def test_source_and_review_foreign_keys(self):
        sources = read('data/traits/sources.csv')
        evidence = read('data/traits/evidence.csv')
        reviews = read('data/traits/review_decisions.csv')
        values = read('data/traits/trait_values.csv')
        conflicts = read('data/traits/conflicts.csv')
        for frame, key in [(sources, 'source_id'), (evidence, 'evidence_id'), (reviews, 'review_decision_id')]:
            self.assertFalse(frame[key].duplicated().any())
            self.assertNotIn('', set(frame[key]))
        for series in [evidence.source_id, reviews.source_id, values.source_id_used, conflicts.source_ids]:
            self.assertFalse(identifiers(series) - set(sources.source_id))
        for series in [reviews.evidence_id, conflicts.evidence_ids]:
            self.assertFalse(identifiers(series) - set(evidence.evidence_id))
        for series in [values.review_decision_ids_used, conflicts.review_decision_ids]:
            self.assertFalse(identifiers(series) - set(reviews.review_decision_id))

    def test_historical_components_match_totals(self):
        table = read('data/historical/beta_summary.csv')
        self.assertEqual(len(table), 18)
        np.testing.assert_allclose(table.mean_component_a + table.mean_component_b, table.mean_total, atol=1e-10)
        np.testing.assert_allclose(table.share_component_a + table.share_component_b, 1, atol=1e-10)

    def test_synthetic_pairing_and_keys(self):
        meta = read('data/synthetic/metadata.csv')
        observations = read('data/synthetic/observations.csv')
        mapping = read('data/synthetic/mapping.csv')
        self.assertEqual(len(meta), 24)
        self.assertTrue(meta.sample_id.str.startswith('DEMO_').all())
        self.assertEqual(meta.groupby('site_id').season.nunique().tolist(), [2] * 12)
        self.assertFalse(observations.duplicated(['sample_id', 'record_id']).any())
        self.assertEqual(set(observations.sample_id), set(meta.sample_id))
        self.assertEqual(set(observations.record_id), set(mapping.record_id))
        self.assertEqual(len(mapping), 47)

    def test_workbook_matches_csv_reading_order(self):
        import json
        catalog = json.loads((ROOT / 'provenance/table_catalog.json').read_text())
        paths = [item['path'] for item in catalog] + [
            'data/synthetic/mapping.csv', 'data/synthetic/observations.csv',
            'data/synthetic/metadata.csv', 'outputs/demo/community_long.csv',
            'outputs/demo/beta_summary.csv', 'outputs/demo/season_tests.csv',
        ]
        book = load_workbook(ROOT / 'data_catalog.xlsx', read_only=True, data_only=True)
        self.assertEqual(len(book.worksheets), len(paths))
        for number, (sheet, path) in enumerate(zip(book.worksheets, paths), 1):
            frame = read(path)
            self.assertTrue(sheet.title.startswith(f'{number:02d}_'))
            self.assertEqual((sheet.max_row - 1, sheet.max_column), frame.shape)
            rows = sheet.iter_rows(values_only=True)
            self.assertEqual(list(next(rows)), list(frame.columns))
            for actual, expected in zip(rows, frame.itertuples(index=False, name=None)):
                for a, b in zip(actual, expected):
                    if b == '':
                        self.assertIsNone(a)
                    elif isinstance(b, (int, float)):
                        self.assertAlmostEqual(a, b, places=12)
                    else:
                        self.assertEqual(a, b)
        book.close()


if __name__ == '__main__':
    unittest.main()
