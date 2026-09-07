"""按阅读顺序生成所有 CSV 的中文目录、预览与 XLSX；不读取外部工作区。"""
from pathlib import Path
import html
import json
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ROOT=Path(__file__).resolve().parents[1]
MEANINGS={
 'review_batch':'审阅批次日期，区分两轮证据处理',
 'taxon_id':'分类单元唯一编号','accepted_or_operational_name':'接受名或操作分类单元名','kingdom':'界','phylum':'门','class':'纲','order':'目','family':'科','genus':'属',
 'output_column':'标准性状编码列名','standard_trait_group_cn':'生物学性状组','standard_trait_id':'标准性状编号','standard_modality_id':'性状模态编号','standard_name_cn':'编码列中文名称','final_data_type':'最终编码数据类型','allowed_values':'允许取值','standardized_value_meaning_zh':'标准值中文含义','missing_value_rule':'缺失与不确定值处理规则',
 'value':'最终分析值','previous_value':'上轮已有值','value_origin':'值来源类别','value_origin_zh':'值来源中文解释','source_id_used':'最终使用的来源编号','candidate_ids_used':'采用的候选编号','review_decision_ids_used':'采用的审阅决定编号',
 'source_id':'来源唯一编号','source_name':'来源名称','citation':'文献引文','citation_or_label':'引文或来源标签','doi_or_url':'DOI 或来源网页','source_scope':'来源区域范围','rights_note_zh':'材料归属与许可说明',
 'evidence_id':'证据编号','evidence_ids':'证据编号集合','accessed_at':'记录的访问日期','source_taxon_name':'来源中的分类单元名称','match_resolution':'分类匹配层级','match_resolution_code':'分类匹配层级代码','authenticity_status':'历史来源真实性标记','evidence_detail_zh':'证据摘要',
 'review_decision_id':'审阅决定唯一编号','candidate_id':'候选值编号','standard_value':'候选标准值','final_review_decision':'最终审阅决定','merge_eligibility_before_conflict_check':'冲突检查前的入库资格','decision_basis_zh':'审阅依据',
 'merge_cell_id':'合并单元格编号','old_value':'原单元格值','cell_resolution_status_zh':'冲突处理状态','accepted_values':'已接受但彼此矛盾的值','candidate_ids':'候选编号集合','review_decision_ids':'审阅决定编号集合','source_ids':'来源编号集合',
 'bar_id':'汇总条唯一编号','data_type':'群落数据类型','data_type_zh':'数据类型中文名','season_zh':'季节中文名','dimension':'多样性维度代码','dimension_zh':'维度中文名','framework':'距离与分解框架','component_a':'第一组分的准确名称','component_b':'第二组分的准确名称','n_pairs':'汇总中的样点对数量','mean_total':'总差异均值','mean_component_a':'第一组分均值','mean_component_b':'第二组分均值','share_component_a':'第一组分在两组分和中的比例','share_component_b':'第二组分比例',
 'effect_scope_zh':'检验范围与置换设计','matrix_label_zh':'响应矩阵中文名称','pseudo_F':'置换多元方差分析伪 F','R2':'解释率','PERMANOVA_FDR':'历史季节检验的校正 p','PERMDISP_FDR':'历史离散度检验的校正 p',
 'selected_environment_variables':'筛选后环境变量','selected_space_variables':'空间变量','n_environment_variables':'环境变量数','n_space_variables':'空间变量数','vpa_panel':'历史 VPA 图面板编号','response_matrix_zh':'响应距离矩阵','model_scope_zh':'模型变量组合','predictors_zh':'模型解释变量','adjusted_R2':'调整后的解释率','permutation_p':'置换检验 p 值',
 'record_id':'虚构原始记录编号','synthetic_name':'虚构记录名称','action':'演示处理动作','reason_zh':'处理理由','sample_id':'虚构样本编号','site_id':'虚构固定点编号','season':'季节代码','data_origin':'数据来源类别','density':'合成密度值（示例单位为个体每平方米）','biomass':'合成生物量值（示例单位为克每平方米）','incidence':'有无值，出现为 1，否则为 0',
 'coefficient':'距离系数，数量型名称见下方术语表','partition':'分解框架','lingoes_constant':'Lingoes 欧氏化的校正常数','n_permutations':'实际枚举的标签排列数','p_permanova':'演示季节检验原始 p','p_permdisp':'演示离散度检验原始 p','q_permanova':'九个响应内 BH 校正的季节 p','q_permdisp':'九个响应内 BH 校正的离散度 p',
}


def main():
    items=json.loads((ROOT/'provenance/table_catalog.json').read_text())
    extra=[
      ('data/synthetic/mapping.csv','合成演示：分类映射','从零构造 47 条记录与目标单元的关系，不是原始名录映射','record_id；taxon_id'),
      ('data/synthetic/observations.csv','合成演示：原始观测','独立随机种子；未拟合真实群落','sample_id + record_id'),
      ('data/synthetic/metadata.csv','合成演示：配对信息','24 个虚构样本；每个固定点对应两季','sample_id；site_id'),
      ('outputs/demo/community_long.csv','合成输出：规则合并后的群落','合成观测经映射，数量求和；有无重新二值化','sample_id + taxon_id'),
      ('outputs/demo/beta_summary.csv','合成输出：分季节 β 汇总','公开演示按三类数据、三个维度计算；不用于支持论文结论','data_type + dimension + season'),
      ('outputs/demo/season_tests.csv','合成输出：配对季节检验','九个总距离矩阵；Lingoes 校正后，固定点内精确枚举 4096 次','data_type + dimension'),
    ]
    for path,title,source,keys in extra:items.append({'path':path,'title_zh':title,'source_zh':source,'keys_zh':keys})
    actual={str(p.relative_to(ROOT)) for folder in ['data','outputs'] for p in (ROOT/folder).rglob('*.csv')}
    assert actual=={x['path'] for x in items},'CSV 目录与实际输出范围不一致'
    schema=pd.read_csv(ROOT/'data/traits/trait_schema.csv').set_index('output_column')
    parts=[]
    with pd.ExcelWriter(ROOT/'data_catalog.xlsx',engine='openpyxl') as writer:
        for i,item in enumerate(items,1):
            df=pd.read_csv(ROOT/item['path'],keep_default_na=False)
            meanings={c:MEANINGS.get(c, schema.loc[c,'standard_name_cn'] if c in schema.index else '') for c in df}
            missing=[c for c,v in meanings.items() if not v]
            if missing:raise ValueError('缺少中文列解释：'+str(missing))
            if item['path'].endswith('trait_matrix.csv'):
                selected=['taxon_id']+list(df.columns[1:7]);preview=df[selected].head(6)
            else:preview=df.head(5)
            if preview.shape[1]>9:preview=preview.iloc[:,:9]
            preview=preview.rename(columns={c:c+'（'+meanings[c]+'）' for c in preview})
            dictionary=pd.DataFrame({'英文列名':list(meanings),'中文解释':list(meanings.values())})
            parts.append(f'<section id="table-{i}"><h2>{i:02d} · {html.escape(item["title_zh"])}</h2><p><a href="../{item["path"]}">{item["path"]}</a> · {len(df)} 行 × {len(df.columns)} 列</p><p>来源：{html.escape(item["source_zh"])}<br>主键／连接键：{html.escape(item["keys_zh"])}</p><div class="table">{preview.to_html(index=False,border=0,escape=True)}</div><p>预览至多 5–6 行、9 列，完整数据在 CSV 和同名顺序工作表中。</p><details><summary>查看全部英文列名与中文解释</summary>{dictionary.to_html(index=False,border=0,escape=True)}</details></section>')
            sheet=f'{i:02d}_'+item['title_zh'].replace('：','_')[:24]
            df.to_excel(writer,sheet_name=sheet,index=False)
            ws=writer.book[sheet];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.fill=PatternFill('solid',fgColor='1B625D');cell.font=Font(color='FFFFFF',bold=True)
                cell.alignment=Alignment(wrap_text=True,vertical='top')
            for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=25
    nav='<nav><a href="../index.html">项目入口</a><a href="../data_catalog.xlsx">下载 XLSX</a></nav>'
    toc=''.join(f'<li><a href="#table-{i}">{i:02d} {html.escape(x["title_zh"])}</a></li>' for i,x in enumerate(items,1))
    intro='''<section><h2>表格如何关联</h2><p>真实性状材料先读 taxa（分类单元），再读 trait_schema（编码定义）和 trait_matrix（分析宽表）；trait_values（来源长表）用 taxon_id + output_column 与宽表对应，再经来源／候选／审阅编号连接 sources、evidence、review_decisions 与 conflicts。分号分隔的编号集合需先拆分再连接。</p><p>trait_values 中 previous_merge_run 表示沿用更早版本；review_decisions 合并两轮共 136 条记录，用 review_batch 区分。FRD 编号来自 5 月 5 日，RGD 编号来自 5 月 15 日；前一批次未保存独立 candidate_id／evidence_id 的字段留空，不补造编号。不是所有沿用值都有审阅编号，但已填写的审阅编号均能连接。性状缺失值是未知，不可视为 0。</p><p>历史统计表仅包含聚合量；不会提供样点观测或两两矩阵。合成观测与合成映射以 record_id 连接，合并输出用 sample_id + taxon_id，配对由 metadata 的 site_id 定义。真实与合成表不能混合用于科学推断。</p><details><summary>常见代码的中文含义</summary><p>incidence：有无；density：密度；biomass：生物量。TD：分类学；FD：功能；PD_proxy：分类树代理。summer／autumn：夏／秋。sorensen 在有无输入下指 Sørensen，数量型对应 Bray–Curtis；jaccard 在有无输入下指 Jaccard。turnover：周转；nestedness_resultant：嵌套结果；replacement：替换；richness_difference：丰富度差异；abundance_difference／abundance_gradient：数量差异／数量梯度；balanced_variation：平衡数量变化。</p><p>GO：开放属级记录；EG：精确属匹配；OS：原始种上卷；SG：同属其他种。accepted_after_user_review：人工审阅接受；eligible_reviewed：已审阅但仍需检查冲突。clear：历史来源标记为清楚，不代表本次独立重查全部原文。</p></details></section>'''
    style='body{font:16px/1.8 "Noto Sans CJK SC","Microsoft YaHei",sans-serif;background:#f2f5ed;color:#213b3d;margin:0}main{max-width:1100px;margin:auto;padding:24px}section{background:#fffef9;border:1px solid #dbe5d9;padding:22px;margin:20px 0;border-radius:12px;scroll-margin-top:70px}nav{position:sticky;top:0;background:#e2edde;padding:15px;display:flex;gap:30px}a{color:#176b62}.table{overflow:auto}table{border-collapse:collapse;font-size:13px;width:100%}td,th{padding:9px;border-bottom:1px solid #dbe5d9;min-width:130px;text-align:left;vertical-align:top}th{background:#e7efe1}summary{cursor:pointer;font-weight:bold}details{margin:12px 0;padding:10px;border-top:1px solid #dbe5d9}p{overflow-wrap:anywhere}'
    page=f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>数据阅读目录</title><style>{style}</style></head><body>{nav}<main><h1>数据目录 · 每张 CSV 的来源、连接键与中文列解释</h1>{intro}<section><details open><summary>推荐阅读顺序（与 XLSX 工作表一致）</summary><ol>{toc}</ol></details></section>{"".join(parts)}</main></body></html>'
    (ROOT/'docs/data_catalog.html').write_text(page,encoding='utf-8')
    print(f'数据交付完成：{len(items)} 张 CSV，逐表中文预览和列解释，XLSX 工作表顺序一致。')


if __name__=='__main__':main()
